from __future__ import annotations

from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "reports" / "10-step-on-people_탐구보고서.docx"

# Design preset: narrative_proposal.
# Named overrides: Malgun Gothic for East Asian glyphs; editorial-cover title
# treatment; compact 9–10 pt table text where numerical evidence needs it.
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
NAVY = "203748"
GOLD = "A17619"
INK = "20252B"
MUTED = "666D75"
LIGHT_FILL = "F4F6F9"
PALE_BLUE = "E8EEF5"
WHITE = "FFFFFF"
GRID = "B8C2CC"


def set_run_font(
    run,
    *,
    size: float | None = None,
    bold: bool | None = None,
    italic: bool | None = None,
    color: str | None = None,
    latin: str = "Calibri",
    east_asia: str = "Malgun Gothic",
) -> None:
    run.font.name = latin
    rpr = run._r.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)
    rfonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.333

    heading_tokens = {
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 12, 6),
        "Heading 3": (12, DARK_BLUE, 8, 4),
    }
    for name, (size, color, before, after) in heading_tokens.items():
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("페이지 ")
    set_run_font(run, size=9, color=MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, instr, separate, text, end):
        run._r.append(node)


def configure_section(section) -> None:
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True

    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("10-step-on-people · 한국어 1.58-bit 언어 모델 탐구")
    set_run_font(r, size=8.5, color=MUTED)

    footer = section.footer
    p = footer.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    add_page_number(p)

    first_header = section.first_page_header
    first_header.paragraphs[0].text = ""
    first_footer = section.first_page_footer
    first_footer.paragraphs[0].text = ""


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = borders.find(qn(f"w:{edge}"))
        if border is None:
            border = OxmlElement(f"w:{edge}")
            borders.append(border)
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), GRID)


def set_table_geometry(table, widths_dxa: list[int]) -> None:
    if sum(widths_dxa) != 9360:
        raise ValueError(f"table widths must sum to 9360 DXA, got {sum(widths_dxa)}")
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr

    width = tbl_pr.find(qn("w:tblW"))
    if width is None:
        width = OxmlElement("w:tblW")
        tbl_pr.append(width)
    width.set(qn("w:w"), "9360")
    width.set(qn("w:type"), "dxa")

    indent = tbl_pr.find(qn("w:tblInd"))
    if indent is None:
        indent = OxmlElement("w:tblInd")
        tbl_pr.append(indent)
    indent.set(qn("w:w"), "120")
    indent.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for value in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(value))
        grid.append(col)

    for row in table.rows:
        for index, (cell, value) in enumerate(zip(row.cells, widths_dxa)):
            cell.width = Inches(value / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(value))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_table_borders(table)


def format_cell(cell, *, header: bool = False, size: float = 9.5) -> None:
    if header:
        set_cell_shading(cell, LIGHT_FILL)
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(2)
        paragraph.paragraph_format.line_spacing = 1.08
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for run in paragraph.runs:
            set_run_font(run, size=size, bold=header, color=INK)


def add_table(doc, headers: list[str], rows: list[list[str]], widths_dxa: list[int], *, size=9.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for index, text in enumerate(headers):
        table.rows[0].cells[index].text = text
    for row_data in rows:
        cells = table.add_row().cells
        for index, text in enumerate(row_data):
            cells[index].text = text
    set_table_geometry(table, widths_dxa)
    for cell in table.rows[0].cells:
        format_cell(cell, header=True, size=size)
    for row in table.rows[1:]:
        for cell in row.cells:
            format_cell(cell, size=size)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_callout(doc, title: str, body: str, *, fill: str = PALE_BLUE) -> None:
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    set_run_font(r, size=10.5, bold=True, color=DARK_BLUE)
    p = cell.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(body)
    set_run_font(r, size=10, color=INK)
    set_table_geometry(table, [9360])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_body(doc, text: str, *, bold_lead: str | None = None) -> None:
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        r = p.add_run(bold_lead)
        set_run_font(r, bold=True, color=DARK_BLUE)
        r = p.add_run(text[len(bold_lead) :])
        set_run_font(r, color=INK)
    else:
        r = p.add_run(text)
        set_run_font(r, color=INK)


def add_heading(doc, text: str, level: int = 1) -> None:
    p = doc.add_paragraph(style=f"Heading {level}")
    r = p.add_run(text)
    set_run_font(r, size={1: 16, 2: 13, 3: 12}[level], bold=True, color={1: BLUE, 2: BLUE, 3: DARK_BLUE}[level])


def add_page_break(doc) -> None:
    doc.add_page_break()


def add_cover(doc) -> None:
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(72)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(18)
    r = p.add_run("정보·인공지능 주제 탐구 보고서")
    set_run_font(r, size=11, bold=True, color=GOLD)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("10-step-on-people")
    set_run_font(r, size=30, bold=True, color=NAVY)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("한국어 1.58-bit 언어 모델을 직접 만들고 비교하기")
    set_run_font(r, size=15, color=DARK_BLUE)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(54)
    r = p.add_run("일반 정밀도 모델과 BitNet b1.58 모델의 직접 구현·학습·비교")
    set_run_font(r, size=10.5, italic=True, color=MUTED)

    add_table(
        doc,
        ["작성 정보", "내용"],
        [
            ["학교", "____________________________"],
            ["학년 · 반", "________학년  ________반"],
            ["학번", "____________________________"],
            ["이름", "____________________________"],
            ["과목", "____________________________"],
            ["담당 교사", "____________________________"],
            ["제출일", "2026년  ______월  ______일"],
        ],
        [2700, 6660],
        size=10,
    )

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(18)
    r = p.add_run("탐구 분야: 저비트 인공지능 · 자연어 처리 · 한국어 생성 모델")
    set_run_font(r, size=9.5, color=MUTED)


def build_report() -> Document:
    doc = Document()
    configure_styles(doc)
    for section in doc.sections:
        configure_section(section)

    add_cover(doc)
    add_page_break(doc)

    add_heading(doc, "1. 탐구 개요")
    add_callout(
        doc,
        "결론부터 말하면",
        "약 1,901만 개의 매개변수(parameter)를 가진 한국어 문장 이어쓰기 모델 두 개를 처음부터 학습하였다. "
        "BitNet b1.58 모델은 모든 선형층을 삼진 가중치 {-1, 0, 1}로 가짜 양자화했고, "
        "동일 구조의 일반 정밀도 기준 모델(FP baseline)과 같은 1,500 최적화 단계로 비교하였다. 두 모델 모두 학습 전보다 검증 손실이 감소하여 구현과 학습이 실제로 작동했음을 확인하였다.",
    )
    add_heading(doc, "1.1 탐구 동기", 2)
    add_body(
        doc,
        "최근 대규모 언어 모델은 자연스러운 글을 생성하지만, 매개변수 수와 연산량이 커 개인용 그래픽카드에서 직접 학습하기 어렵다. "
        "그래서 ‘가중치를 아주 적은 비트로 표현하면 작은 GPU에서도 언어 모델의 원리를 실험할 수 있을까?’라는 질문에서 탐구를 시작했다. "
        "완성된 챗봇이나 지식 질의응답을 만드는 대신, 문장의 앞부분을 주면 다음 토큰(token)을 이어 쓰는 가장 기본적인 언어 모델 기능에 범위를 한정하였다."
    )
    add_heading(doc, "1.2 탐구 목적과 탐구 문제", 2)
    add_body(
        doc,
        "본 탐구의 목적은 제한된 학교·가정용 하드웨어에서도 저비트 언어 모델의 핵심 원리를 직접 구현하고, 동일 조건의 일반 정밀도 모델과 비교하여 장점과 한계를 수치로 확인하는 데 있다."
    )
    add_table(
        doc,
        ["탐구 질문", "확인 방법"],
        [
            ["1.58-bit 방식이 실제 코드와 학습 과정으로 구현되는가?", "삼진 코드 분포, 기울기(gradient), 체크포인트를 검사한다."],
            ["제한된 2GB GPU에서 두 모델을 공정하게 학습할 수 있는가?", "같은 구조, 데이터 분할, 최적화 단계, 유효 배치 크기를 사용한다."],
            ["학습 전보다 다음 토큰 예측이 개선되는가?", "학습 전과 학습 후의 검증 손실 및 혼란도(perplexity)를 비교한다."],
            ["저비트 구현이 곧바로 더 빠르고 작은가?", "체크포인트 크기와 실제 생성 시간을 함께 기록한다."],
        ],
        [3600, 5760],
        size=9.0,
    )
    add_body(
        doc,
        "이 탐구의 성공 기준은 ‘사람처럼 대화하는 모델’을 완성하는 것이 아니다. 동일 조건의 두 모델을 직접 구현하고 학습하여 수치 자료와 실행 가능한 시연 프로그램을 남기는 것을 성공 기준으로 정하였다."
    )
    add_callout(
        doc,
        "프로젝트 이름에 대한 설명",
        "‘10-step-on-people’의 ‘10-step’은 프로젝트 이름이며 실제 학습 단계 수를 뜻하지 않는다. 본 실험의 실제 학습량은 모델별 1,500 최적화 단계이다.",
        fill=LIGHT_FILL,
    )
    add_heading(doc, "2. 이론적 배경")
    add_heading(doc, "2.1 자기회귀 언어 모델", 2)
    add_body(
        doc,
        "자기회귀 언어 모델은 앞에서 본 토큰들을 이용해 다음 토큰의 확률을 예측한다. 예를 들어 ‘오늘 학교에서’를 입력하면 모델은 그 뒤에 올 가능성이 높은 토큰을 하나씩 뽑아 문장을 이어 간다. "
        "학습에서는 입력 토큰 열과 한 칸 오른쪽으로 이동한 정답 열을 비교하고, 정답 확률이 높아지도록 교차 엔트로피 손실(cross-entropy loss)을 줄인다."
    )
    add_callout(
        doc,
        "다음 토큰 예측",
        "입력: x₁, x₂, …, xₜ₋₁  →  예측: P(xₜ | x₁, x₂, …, xₜ₋₁)  →  새 토큰을 입력 뒤에 붙여 반복",
        fill=LIGHT_FILL,
    )
    add_heading(doc, "2.2 일반 정밀도 기준 모델과 BitNet b1.58", 2)
    add_body(
        doc,
        "FP baseline은 PyTorch의 일반 선형층을 사용한다. BitNet 모델은 attention, feed-forward network, 출력 head의 선형층을 BitLinear로 바꾼다. "
        "BitLinear는 학습 가능한 full-precision master weight를 유지하지만 forward 계산에서는 가중치의 절댓값 평균을 scale로 삼아 {-1, 0, 1} 가운데 하나로 바꾼다. 세 상태를 구분하는 정보량 log₂(3)이 약 1.58이어서 b1.58이라 부른다."
    )
    add_body(
        doc,
        "활성값은 token마다 절댓값 최댓값을 기준으로 int8 범위에 가짜 양자화했다. 반올림 연산은 그대로는 gradient가 끊기므로, backward에서는 항등함수처럼 gradient를 흘리는 straight-through estimator(STE)를 사용했다. embedding과 normalization은 full precision으로 남겼다."
    )
    add_heading(doc, "2.3 가짜 양자화의 의미", 2)
    add_body(
        doc,
        "이 구현은 저비트 계산 방식을 학습 과정에서 흉내 내는 fake quantization이다. 실제 checkpoint에는 master weight가 float 형태로 저장되고, 전용 packed 형식이나 저비트 CUDA kernel을 사용하지 않는다. 따라서 ‘1.58-bit 원리를 구현했다’고 말할 수는 있지만 파일 크기, 속도, 전력 절감을 달성했다고 주장할 수는 없다."
    )
    add_heading(doc, "3. 탐구 방법")
    add_heading(doc, "3.1 데이터와 토크나이저", 2)
    add_body(
        doc,
        "Hugging Face의 oz1115/korean-pretraining-corpus-ko에서 revision을 고정하고 val.pkl만 내려받아 교육용 subset으로 다시 분할했다. 원래 validation subset을 학습에 재사용한 것이므로 일반적인 연구 평가용 split과는 다르다. "
        "pickle은 임의 class 생성을 금지한 restricted unpickler로 읽은 뒤 shape와 token 범위를 검증했고, uint16 NumPy 배열로 변환하여 memory-map으로 사용했다."
    )
    add_table(
        doc,
        ["항목", "설정", "이유"],
        [
            ["전체 token", "12,490,752", "691MB 전체 자료 대신 빠른 실험 범위 확보"],
            ["행 구조", "24,396 × 512", "문서 경계를 넘지 않는 학습 sample 생성"],
            ["split", "row 기준 90% / 10%, seed 42", "두 모델에 동일하고 비중복인 train/validation 제공"],
            ["tokenizer", "BPE, vocab 32,000", "korean-gpt-150m-ko tokenizer revision 고정"],
            ["context", "128 tokens", "2GB VRAM에서 학습 가능하도록 제한"],
        ],
        [1900, 2700, 4760],
        size=9.0,
    )
    add_heading(doc, "3.2 모델 구조와 공정한 비교 조건", 2)
    add_table(
        doc,
        ["구성", "공통값"],
        [
            ["Transformer", "decoder 4 layers, hidden 256, 8 heads, FFN 768"],
            ["연산 구성", "RoPE, RMSNorm, ReLU², dropout 0.1, bias 없음"],
            ["parameter", "두 모델 모두 19,007,744개"],
            ["차이", "baseline은 nn.Linear, BitNet은 attention·FFN·LM head에 BitLinear"],
            ["optimizer", "AdamW, LR 3e-4, warm-up 150, cosine decay, gradient clip 1.0"],
            ["학습량", "두 모델 모두 1,500 optimizer steps"],
        ],
        [2300, 7060],
    )
    add_heading(doc, "3.3 탐구 절차", 2)
    add_table(
        doc,
        ["단계", "수행 내용", "확인 자료"],
        [
            ["1. 자료 준비", "고정된 revision의 val.pkl과 토크나이저를 내려받고 안전하게 변환", "shape·토큰 범위·분할 검증 결과"],
            ["2. 모델 구현", "공통 Transformer와 BitLinear를 구현하고 단위 테스트 수행", "삼진 코드·int8 범위·기울기 테스트"],
            ["3. GPU 사전 점검", "실제 순전파·역전파로 가능한 micro-batch를 측정", "run_profile.json"],
            ["4. BitNet 학습", "BitNet을 최대 1,500 단계까지 먼저 학습", "checkpoint·metric JSONL"],
            ["5. 기준 모델 학습", "동일한 단계와 배치 조건으로 FP baseline 학습", "checkpoint·metric JSONL"],
            ["6. 결과 비교", "고정 문장 4개 생성, 손실·혼란도·시간 비교", "results.json·시연 프로그램"],
        ],
        [1450, 5350, 2560],
        size=9.0,
    )
    add_heading(doc, "4. 시행 중 부딪힌 문제와 해결한 방법")
    add_table(
        doc,
        ["문제", "원인", "해결 방법", "결과"],
        [
            ["전체 말뭉치 준비가 오래 걸림", "원본 전체 다운로드와 변환 비용이 큼", "34.7MB val.pkl만 고정 revision으로 받아 교육용 subset으로 재분할", "약 1,249만 token으로 빠른 실험 가능"],
            ["pickle 보안 위험", "일반 pickle은 객체 생성 코드를 포함할 수 있음", "class loading을 금지한 restricted unpickler와 shape·범위 검증 적용", "허용된 정수 token 배열만 통과"],
            ["MX450 전용 VRAM 2GB 제한", "batch가 크면 CUDA OOM 가능", "BitNet 실제 forward/backward/optimizer preflight 후 micro-batch와 accumulation 자동 조절", "micro-batch 8, accumulation 4, effective batch 32"],
            ["처음 설치한 PyTorch가 CPU용", "CUDA wheel이 아닌 환경", "CUDA 12.6 wheel로 교체하고 장치 이름·가용성을 확인", "MX450에서 실제 학습 수행"],
            ["1.58-bit인데 파일이 작아지지 않음", "master weight를 float로 저장하고 packed 저장 미구현", "fake quantization의 범위를 문서화하고 크기 절감 주장을 제외", "두 checkpoint 모두 약 72.52MB"],
            ["BitNet 생성이 더 느림", "일반 PyTorch에서 매 forward마다 양자화를 흉내 냄", "전용 kernel이 없다는 한계를 명시하고 latency를 그대로 공개", "BitNet 약 1.05–1.53초, baseline 약 0.40–0.43초"],
            ["생성 문장이 매끄럽지 않음", "작은 모델·짧은 학습·위키 문법이 섞인 data", "고정 prompt와 seed로 재현 가능한 비교를 만들고 품질 부족을 숨기지 않음", "네 prompt 모두 생성은 완료했으나 문장 품질은 제한적"],
        ],
        [1800, 2000, 3300, 2260],
        size=8.4,
    )
    add_callout(
        doc,
        "공정성을 지킨 결정",
        "후속 preflight에서는 micro-batch 16도 성공했지만, 완료된 본 비교 실험은 두 모델 모두 micro-batch 8과 accumulation 4를 유지했다. 한 모델만 더 유리한 batch로 다시 평가하지 않았다.",
        fill=LIGHT_FILL,
    )
    add_heading(doc, "5. 탐구 결과")
    add_heading(doc, "5.1 정량 결과", 2)
    add_table(
        doc,
        ["모델", "step 0 loss", "best loss", "개선량", "perplexity", "학습 시간"],
        [
            ["BitNet b1.58", "10.3986", "6.6938", "3.7048", "807.41", "31분 16초"],
            ["FP baseline", "10.4161", "6.5924", "3.8238", "729.49", "28분 22초"],
        ],
        [2100, 1350, 1350, 1250, 1500, 1810],
        size=9.2,
    )
    add_body(
        doc,
        "두 모델 모두 step 0보다 best validation loss가 크게 낮아졌다. 이는 dataset 준비, causal mask, backward, optimizer, checkpoint 저장으로 이어지는 학습 파이프라인이 실제로 작동했다는 직접적인 증거다. "
        "이번 조건에서는 FP baseline의 loss와 perplexity가 BitNet보다 조금 더 낮았다. 따라서 BitNet이 더 정확했다고 결론내릴 수는 없다."
    )
    add_heading(doc, "5.2 BitNet 내부와 실행 결과", 2)
    add_table(
        doc,
        ["증거 항목", "측정값", "해석"],
        [
            ["최종 ternary 0 비율", "31.86%", "BitLinear weight code의 약 3분의 1이 0 상태"],
            ["linear ternary 대상 weight", "10,813,440개", "embedding·normalization을 제외한 양자화 대상"],
            ["preflight peak allocated VRAM", "541.31MB", "micro-batch 8 실제 optimizer step에서 측정"],
            ["checkpoint 크기", "BitNet/FP 모두 약 72.52MB", "fake quantization이라 저장 공간 절감 없음"],
        ],
        [2800, 1900, 4660],
    )
    add_heading(doc, "5.3 문장 이어쓰기 예시", 2)
    add_body(
        doc,
        "고정 prompt ‘오늘 학교에서’에 대해 BitNet은 ‘다시 [[카(마블 시리즈)|그트]]의 …’처럼 위키 문법이 섞인 조각을 만들었고, baseline은 ‘, [[제2차]]에 의하면 해당에서 …’처럼 이어 썼다. "
        "둘 다 32 token 생성을 완료했지만 자연스러운 학교생활 문장에는 이르지 못했다. 이 결과는 데모가 작동한다는 증거이면서 동시에 현재 품질의 한계를 보여 준다."
    )
    add_callout(
        doc,
        "결과 해석의 핵심",
        "‘저비트 모델이 무조건 더 좋다’가 아니라, 2GB GPU에서 1.58-bit 학습 원리를 직접 구현하고 동일 조건의 FP 모델과 수치로 비교했다.",
    )
    add_heading(doc, "6. 한계")
    add_body(doc, "첫째, 모델이 약 1,901만 parameter로 매우 작고 1,500 step만 학습되어 자연스러운 한국어 생성 능력이 부족하다. 결과 문장에 위키 링크 문법과 불완전한 표현이 나타났다.")
    add_body(doc, "둘째, 원래 validation 용도였던 subset을 교육용 train/validation으로 다시 나누었다. 원 자료의 상세한 수집·정제 설명도 제한적이므로, 일반화 성능이나 data 품질을 강하게 주장할 수 없다.")
    add_body(doc, "셋째, 가중치와 활성값은 fake quantization이며 full-precision master weight를 유지한다. 실제 1.58-bit packing과 전용 kernel이 없어서 checkpoint 크기, 추론 속도, 전력 소비의 이점을 검증하지 못했다.")
    add_body(doc, "넷째, 평가는 validation loss·perplexity와 고정 prompt 4개에 한정되었다. 사람 평가, 문법성 평가, 다양한 sampling 설정 비교가 없고 seed도 42 하나만 사용했다.")
    add_body(doc, "다섯째, 한 대의 MX450에서만 실행했으므로 다른 GPU나 CPU에서의 속도와 메모리 사용량을 대표하지 않는다.")
    add_heading(doc, "7. 결론 및 느낀 점")
    add_heading(doc, "7.1 결론", 2)
    add_body(
        doc,
        "이번 탐구에서는 한국어 token data를 안전하게 준비하고, decoder-only Transformer와 BitLinear를 직접 구현하여, FP baseline과 BitNet b1.58을 같은 조건에서 학습했다. "
        "두 모델 모두 검증 손실이 감소했고 checkpoint, metric JSONL, GPU profile, 고정 prompt 출력이 남아 있어 실험 수행을 재현하고 확인할 수 있다."
    )
    add_body(
        doc,
        "가장 중요한 배움은 양자화의 이론적 표현 비트 수와 실제 프로그램의 파일 크기·속도가 같지 않다는 점이다. 알고리즘만 흉내 내는 fake quantization 단계와, packed 저장 및 전용 kernel까지 갖춘 실용 시스템 단계는 구분해야 한다. "
        "후속 탐구에서는 data 정제를 강화하고 학습량을 늘린 뒤, 실제 ternary packing 또는 BitNet 전용 kernel을 적용해 메모리와 속도를 다시 측정할 수 있다."
    )
    add_heading(doc, "7.2 탐구를 통해 느낀 점", 2)
    add_body(
        doc,
        "처음에는 1.58-bit라는 이름만 보고 모델 파일이 즉시 작아지고 속도도 빨라질 것으로 예상하였다. 그러나 직접 구현해 보니 학습용 master weight와 전용 연산 kernel의 유무가 실제 성능을 좌우한다는 사실을 알게 되었다. "
        "또한 결과가 기대만큼 자연스럽지 않더라도 손실, 학습 시간, 생성 예시를 그대로 기록하는 과정이 탐구의 신뢰성을 높인다는 점을 배웠다. 제한된 장비에서도 문제를 단계별로 줄이고 측정 자료를 남기면 인공지능의 핵심 원리를 직접 검증할 수 있다는 점이 가장 의미 있었다."
    )
    add_page_break(doc)
    add_heading(doc, "부록 A. 재현 방법과 증거 파일")
    add_heading(doc, "A.1 실행 순서", 2)
    add_table(
        doc,
        ["목적", "명령 또는 방법"],
        [
            ["데모 실행", "Windows에서 RUN_DEMO.bat 더블클릭"],
            ["data/tokenizer 준비", "python -m onebit_llm.prepare"],
            ["두 모델 학습", "python -m onebit_llm.train --model all --config configs/demo.yaml"],
            ["평가 보고서 생성", "python -m onebit_llm.evaluate"],
            ["전체 test", "pytest"],
        ],
        [2500, 6860],
        size=9.2,
    )
    add_heading(doc, "A.2 실험 증거", 2)
    add_table(
        doc,
        ["파일", "내용"],
        [
            ["artifacts/summary.json", "두 모델의 step, 초기/best loss, 시간, parameter 수"],
            ["artifacts/run_profile.json", "GPU, CUDA, batch, accumulation, peak VRAM"],
            ["artifacts/metrics/*.jsonl", "100 step 간격 validation loss와 ternary 0 비율"],
            ["artifacts/checkpoints/*/best.pt", "실제로 데모가 불러오는 학습 checkpoint"],
            ["artifacts/evaluation/results.json", "고정 prompt 4개, latency, perplexity, 생성 결과"],
        ],
        [3800, 5560],
        size=9.2,
    )
    add_heading(doc, "참고 자료", 1)
    refs = [
        "[1] Wang et al., “BitNet: Scaling 1-bit Transformers for Large Language Models,” arXiv:2310.11453, 2023. https://arxiv.org/abs/2310.11453",
        "[2] Ma et al., “The Era of 1-bit LLMs: All Large Language Models are in 1.58 Bits,” arXiv:2402.17764, 2024. https://arxiv.org/abs/2402.17764",
        "[3] oz1115, korean-pretraining-corpus-ko, Hugging Face. 사용 revision: 324488e1befab3a9b4eac7571bf1557ef4a4eeec. https://huggingface.co/datasets/oz1115/korean-pretraining-corpus-ko",
        "[4] oz1115, korean-gpt-150m-ko tokenizer, Hugging Face. 사용 revision: 669560d7d1de8c3213e43a343e577a80ad6cb7ee. https://huggingface.co/oz1115/korean-gpt-150m-ko",
    ]
    for ref in refs:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.first_line_indent = Inches(-0.25)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(ref)
        set_run_font(r, size=8.0, color=INK)

    return doc


def audit_output(path: Path) -> None:
    """Fail fast if the generated document drifts from its design tokens."""
    doc = Document(path)
    section = doc.sections[0]
    assert section.page_width == Inches(8.5)
    assert section.page_height == Inches(11)
    assert section.top_margin == Inches(1)
    assert section.right_margin == Inches(1)
    assert section.bottom_margin == Inches(1)
    assert section.left_margin == Inches(1)
    assert abs(section.header_distance - Inches(0.492)) < 1000
    assert abs(section.footer_distance - Inches(0.492)) < 1000
    assert section.different_first_page_header_footer

    normal = doc.styles["Normal"]
    assert normal.font.name == "Calibri"
    assert normal.font.size == Pt(11)
    assert normal.paragraph_format.space_after == Pt(8)
    assert abs(normal.paragraph_format.line_spacing - 1.333) < 0.001
    for name, size in (("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 12)):
        assert doc.styles[name].font.size == Pt(size)

    for table in doc.tables:
        tbl_pr = table._tbl.tblPr
        assert tbl_pr.find(qn("w:tblW")).get(qn("w:w")) == "9360"
        assert tbl_pr.find(qn("w:tblInd")).get(qn("w:w")) == "120"
        assert tbl_pr.find(qn("w:tblLayout")).get(qn("w:type")) == "fixed"
        grid_widths = [int(col.get(qn("w:w"))) for col in table._tbl.tblGrid]
        assert sum(grid_widths) == 9360
        for row in table.rows:
            cell_widths = [
                int(cell._tc.get_or_add_tcPr().find(qn("w:tcW")).get(qn("w:w")))
                for cell in row.cells
            ]
            assert cell_widths == grid_widths

    with ZipFile(path) as archive:
        page_breaks = archive.read("word/document.xml").count(b'w:type="page"')
    assert page_breaks == 2, f"expected cover and appendix page breaks, found {page_breaks}"


def main() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = build_report()
    doc.save(OUTPUT)
    audit_output(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
